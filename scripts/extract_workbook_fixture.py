"""Extract a golden-test fixture from a Google Sheets export of the contract workbook.

Usage: uv run python scripts/extract_workbook_fixture.py <workbook.xlsx> <out.json>

Reads only the team tabs (never Contact Info). Blocks are located by their
labels, not fixed cells, because the tabs have drifted apart. Values are the
sheet's cached results, so the export must come straight from Google Sheets.
"""

import json
import re
import sys

import openpyxl

TEAM_TAB = re.compile(r"^([A-Z]{2}) \((.+)$")
BUDGET_LABELS = {
    "base": "Base Draft Budget",
    "starting": "Starting",
    "contracts": "Regular Contracts",
    "buyouts": "Bad Contracts",
    "farm": "Farm Contracts",
    "cash_in": "acquired by trade",
    "missed_ip": "Min. IP Penalties",
    "spent": "Draft Budget Spent",
    "remaining": "Remaining",
}
H, K, M, O, Q, S, U, AA = 8, 11, 13, 15, 17, 19, 21, 27  # noqa: E741 (sheet column letters)
B, C, F, G, I, J = 2, 3, 6, 7, 9, 10  # noqa: E741


def text(v):
    if isinstance(v, float) and v.is_integer():
        return int(v)  # Google Sheets exports whole numbers as floats
    return v.strip() if isinstance(v, str) else v


def team(ws) -> dict:
    cell = lambda r, c: text(ws.cell(r, c).value)  # noqa: E731
    out = {"tab": ws.title, "after_season": cell(1, 2), "budget": {}, "contracts": [], "buyouts": [], "farm": []}
    for r in range(1, 40):
        label = cell(r, H)
        for key, needle in BUDGET_LABELS.items():
            if isinstance(label, str) and needle in label and key not in out["budget"]:
                out["budget"][key] = cell(r, K)
    for r in range(1, ws.max_row + 1):
        m = cell(r, M)
        if isinstance(m, str) and m.endswith("player contract -->>") and cell(r, O):
            if None in (cell(r, Q), cell(r, S), cell(r, U)):
                raise ValueError(f"{ws.title}: contract for {cell(r, O)} at row {r} is missing price/year/length")
            out["contracts"].append(
                {
                    "player": cell(r, O),
                    "original_price": cell(r, Q),
                    "year_signed": cell(r, S),
                    "length": cell(r, U),
                    "sheet_price": cell(r, AA),
                }
            )
        if isinstance(m, str) and m.endswith("dropped contract") and cell(r, O):
            drop_row = next((rr for rr in range(r + 1, r + 5) if cell(rr, M) == "Year Dropped"), None)
            if drop_row is None:
                raise ValueError(f"{ws.title}: no 'Year Dropped' row under dropped contract at M{r}")
            out["buyouts"].append(
                {
                    "player": cell(r, O),
                    "original_price": cell(r, Q),
                    "year_signed": cell(r, S),
                    "length": cell(r, U),
                    "sheet_orig_sign": cell(r, AA),
                    "dropped_in": cell(drop_row, O),
                    "sheet_penalty": cell(drop_row, AA),
                }
            )
        b = cell(r, B)
        if isinstance(b, str) and re.fullmatch(r"Farm \d+", b) and cell(r, C):
            out["farm"].append(
                {
                    "player": cell(r, C),
                    "salary": cell(r, F),
                    "drafted": cell(r, G),
                    "mlb": cell(r, I),
                    "keep": cell(r, J),
                }
            )
    return out


def main(src: str, dest: str) -> None:
    wb = openpyxl.load_workbook(src, data_only=True)
    teams = {}
    for name in wb.sheetnames:
        if m := TEAM_TAB.match(name):
            teams[m.group(1)] = team(wb[name])
    with open(dest, "w") as f:
        json.dump({"source": src.rsplit("/", 1)[-1], "teams": teams}, f, indent=1, default=str)
    for code, t in teams.items():
        print(code, len(t["contracts"]), "contracts", len(t["buyouts"]), "buyouts", len(t["farm"]), "farm", t["budget"])


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
